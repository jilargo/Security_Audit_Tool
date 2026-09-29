"""
Professional formatting for the generated Excel audit report.

Everything visual lives here so that core/SecurityOperations.py stays focused
on collecting data. The styler:

* converts the collectors' string values into real numbers and dates so the
  workbook sorts, filters and charts correctly;
* rebuilds the Dashboard as a cover page (title banner, audit details, security
  posture, KPI cards, collector summary, top-process chart);
* turns the raw ``netsh`` firewall dump into a readable, colour-coded table;
* gives every data sheet a consistent header, zebra striping, borders, sensible
  column widths, an autofilter and print settings;
* highlights findings (suspicious processes and PowerShell commands, disabled
  accounts, firewall states, network states).
"""

import re
from datetime import datetime

from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

# --------------------------------------------------------------------------
# Palette
# --------------------------------------------------------------------------
NAVY = "1F3864"
NAVY_SOFT = "2A4A7B"
ACCENT = "2E75B6"
TEXT = "1F2937"
MUTED = "64748B"
BORDER = "D6DEE8"
STRIPE = "F3F7FC"
LABEL_BG = "EEF3FA"

RISK_STYLES = {
    "LOW": ("DCFCE7", "166534"),
    "MEDIUM": ("FEF3C7", "92400E"),
    "HIGH": ("FFEDD5", "9A3412"),
    "CRITICAL": ("FEE2E2", "991B1B"),
}

GREEN = ("DCFCE7", "166534")
RED = ("FEE2E2", "991B1B")
AMBER = ("FEF3C7", "92400E")
NEUTRAL = ("EEF2F6", "475569")

# Tab colours group the worksheets by theme
TAB_COLORS = {
    "Dashboard": NAVY,
    "Employee Info": "4472C4",
    "Windows Defender": "2E7D32",
    "Other Antivirus": "2E7D32",
    "Windows Events": "C62828",
    "Firewall Status": "C62828",
    "Network Connections": "00838F",
    "Powershell History": "6A1B9A",
    "Running Processes": "EF6C00",
    "Scheduled Tasks": "EF6C00",
    "Startup Programs": "EF6C00",
    "Installed Apps List": "5B6B7F",
    "System Information": "5B6B7F",
    "USB History": "00838F",
    "Local Users": "6A1B9A",
}

# Collectors shown on the cover page, in the order they run
COLLECTOR_SHEETS = [
    ("Antivirus Protection", "Windows Defender"),
    ("Firewall Profiles", "Firewall Status"),
    ("Windows Event Logs", "Windows Events"),
    ("Running Processes", "Running Processes"),
    ("Network Connections", "Network Connections"),
    ("PowerShell History", "Powershell History"),
    ("Scheduled Tasks", "Scheduled Tasks"),
    ("Startup Programs", "Startup Programs"),
    ("Installed Applications", "Installed Apps List"),
    ("USB History", "USB History"),
    ("Local User Accounts", "Local Users"),
    ("System Information", "System Information"),
]

# Compared against the executable name with any ".exe" suffix removed
SUSPICIOUS_PROCESSES = frozenset({
    "powershell", "pwsh", "cmd", "mshta", "wscript", "cscript",
    "regsvr32", "rundll32", "certutil", "bitsadmin", "msiexec", "installutil",
})

INTEGER_RE = re.compile(r"^[+-]?\d+$")
DECIMAL_RE = re.compile(r"^[+-]?\d+\.\d+$")
DATE_FORMATS = (
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
    "%a %b %d %H:%M:%S %Y",
    "%b %d %H:%M:%S %Y",
    "%B %d, %Y %I:%M:%S %p",
    "%B %d, %Y %I:%M %p",
    "%m/%d/%Y %H:%M:%S",
    "%d/%m/%Y %H:%M:%S",
)


class ReportStyler:
    """Applies the report design to an openpyxl workbook."""

    # ------------------------------------------------------------------
    # Entry point
    # ------------------------------------------------------------------
    @classmethod
    def apply(cls, wb, employee_data, computer_name, risk_level, firewall_text,
              defender_status, suspicious_processes, top_processes):
        """Restyle a freshly written audit workbook in place."""
        cls.coerce_value_types(wb)

        cls._rebuild_firewall_sheet(wb, firewall_text)
        cls._rebuild_employee_sheet(wb, employee_data)

        # Counted after the rebuilds so the summary reflects the final sheets
        record_counts = {
            name: max(wb[name].max_row - 1, 0)
            for name in wb.sheetnames if name != "Dashboard"
        }

        context = {
            "employee": employee_data,
            "computer": computer_name,
            "risk_level": risk_level,
            "defender": defender_status,
            "suspicious_processes": suspicious_processes,
            "top_processes": top_processes,
            "record_counts": record_counts,
        }

        for name in wb.sheetnames:
            if name == "Dashboard":
                continue
            cls.style_data_sheet(wb[name], name)

        cls.build_dashboard(wb, context)
        cls.apply_highlights(wb)

    # ------------------------------------------------------------------
    # Value types
    # ------------------------------------------------------------------
    @classmethod
    def coerce_value_types(cls, wb):
        """Turn numeric and date-like strings into real Excel values."""
        for name in wb.sheetnames:
            if name == "Dashboard":
                continue
            ws = wb[name]
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.value = cls._coerce(cell.value)

    @staticmethod
    def _coerce(text):
        """Best-effort conversion of a single cell value."""
        stripped = text.strip()
        if not stripped:
            return text

        looks_like_date = len(stripped) >= 8 and any(ch.isdigit() for ch in stripped) \
            and any(ch in stripped for ch in "-,:/")
        if looks_like_date:
            for fmt in DATE_FORMATS:
                try:
                    return datetime.strptime(stripped, fmt)
                except ValueError:
                    continue

        if INTEGER_RE.match(stripped) and not (len(stripped) > 1 and stripped[0] == "0"):
            try:
                return int(stripped)
            except ValueError:
                pass
        elif DECIMAL_RE.match(stripped):
            try:
                return float(stripped)
            except ValueError:
                pass

        return text

    # ------------------------------------------------------------------
    # Data sheets
    # ------------------------------------------------------------------
    @classmethod
    def style_data_sheet(cls, ws, name=""):
        """Apply the shared table design to one worksheet."""
        if not ws or ws.max_column == 0:
            return

        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = TAB_COLORS.get(name, ACCENT)

        max_col = ws.max_column
        max_row = ws.max_row
        cls._humanise_headers(ws)

        # A single-column sheet is a message, not a table
        if max_col == 1:
            cls._style_notice_sheet(ws)
            return

        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(fill_type="solid", fgColor=NAVY)
        side = Side(style="thin", color=BORDER)
        border = Border(left=side, right=side, top=side, bottom=side)

        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center",
                                       wrap_text=True)
        ws.row_dimensions[1].height = 26

        if max_row < 2:
            cls._add_no_records_notice(ws, max_col)
            return

        stripe_fill = PatternFill(fill_type="solid", fgColor=STRIPE)
        wrap_columns = cls._wide_columns(ws, max_col)

        for row_index, row in enumerate(ws.iter_rows(min_row=2), start=2):
            stripe = row_index % 2 == 0
            for cell in row:
                cell.border = border
                cell.font = Font(name="Calibri", size=10, color=TEXT)
                if stripe:
                    cell.fill = stripe_fill
                cell.alignment = cls._data_alignment(cell, cell.column in wrap_columns)
                if isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm:ss"
                elif isinstance(cell.value, bool):
                    pass
                elif isinstance(cell.value, int):
                    cell.number_format = "0"
                elif isinstance(cell.value, float):
                    cell.number_format = "#,##0.##"

        for column_index in range(1, max_col + 1):
            ws.column_dimensions[get_column_letter(column_index)].width = \
                cls._column_width(ws, column_index)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"
        cls._setup_printing(ws)

    @staticmethod
    def _humanise_headers(ws):
        """'first_name' -> 'First Name' for machine-style headers."""
        for cell in ws[1]:
            value = cell.value
            if isinstance(value, str) and ("_" in value or value.islower()):
                cell.value = value.replace("_", " ").title()

    @staticmethod
    def _wide_columns(ws, max_col, threshold=48):
        """Columns long enough to need wrapped text."""
        wide = set()
        for index in range(1, max_col + 1):
            longest = 0
            for row_index in range(2, min(ws.max_row, 200) + 1):
                value = ws.cell(row=row_index, column=index).value
                if isinstance(value, str):
                    longest = max(longest, len(value))
            if longest > threshold:
                wide.add(index)
        return wide

    @staticmethod
    def _column_width(ws, index, minimum=11, maximum=58):
        longest = 0
        for row_index in range(1, min(ws.max_row, 300) + 1):
            value = ws.cell(row=row_index, column=index).value
            if value is not None:
                longest = max(longest, len(str(value)))
        return max(minimum, min(longest + 3, maximum))

    @staticmethod
    def _data_alignment(cell, wrap):
        if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
            return Alignment(horizontal="right", vertical="top")
        if isinstance(cell.value, datetime):
            return Alignment(horizontal="center", vertical="top")
        return Alignment(horizontal="left", vertical="top", wrap_text=wrap)

    @classmethod
    def _style_notice_sheet(cls, ws):
        """Render a one-column sheet as a friendly notice block."""
        message = ws.cell(row=2, column=1).value or "No data collected."
        if ws.max_row > 2:
            ws.delete_rows(2, ws.max_row - 1)
        ws.column_dimensions["A"].width = 80

        ws["A1"].value = "Collector note"
        cls._paint_row(ws, 1, 1, 1,
                       fill=PatternFill(fill_type="solid", fgColor=NAVY),
                       font=Font(name="Calibri", size=11, bold=True, color="FFFFFF"),
                       alignment=Alignment(horizontal="left", vertical="center", indent=1))
        ws.row_dimensions[1].height = 26

        ws["A2"].value = str(message)
        cls._paint_row(ws, 2, 1, 1,
                       fill=PatternFill(fill_type="solid", fgColor=LABEL_BG),
                       font=Font(name="Calibri", size=10, italic=True, color=MUTED),
                       alignment=Alignment(horizontal="left", vertical="center", wrap_text=True,
                                          indent=1))
        ws.row_dimensions[2].height = 30
        cls._setup_printing(ws)

    @classmethod
    def _add_no_records_notice(cls, ws, max_col):
        """Show a note instead of a header-only table."""
        ws.cell(row=2, column=1).value = "No records were collected for this category."
        cls._paint_row(ws, 2, 1, max_col,
                       fill=PatternFill(fill_type="solid", fgColor=LABEL_BG),
                       font=Font(name="Calibri", size=10, italic=True, color=MUTED),
                       alignment=Alignment(horizontal="left", vertical="center", indent=1))
        ws.row_dimensions[2].height = 24

    @staticmethod
    def _setup_printing(ws):
        """Landscape, fit to one page wide, repeat the header row."""
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.print_title_rows = "1:1"
        ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.6,
                                      header=0.3, footer=0.3)
        ws.oddHeader.left.text = "Security Audit Tool"
        ws.oddHeader.left.size = 8
        ws.oddHeader.left.color = MUTED
        ws.oddHeader.right.text = "&A"
        ws.oddHeader.right.size = 8
        ws.oddHeader.right.color = MUTED
        ws.oddFooter.left.text = "Generated &D"
        ws.oddFooter.left.size = 8
        ws.oddFooter.left.color = MUTED
        ws.oddFooter.right.text = "Page &P of &N"
        ws.oddFooter.right.size = 8
        ws.oddFooter.right.color = MUTED

    # ------------------------------------------------------------------
    # Rebuilt sheets
    # ------------------------------------------------------------------
    @classmethod
    def _rebuild_firewall_sheet(cls, wb, firewall_text):
        """Replace the raw netsh dump with a structured profile table."""
        if "Firewall Status" not in wb.sheetnames:
            return

        ws = wb["Firewall Status"]
        position = wb.sheetnames.index("Firewall Status")
        wb.remove(ws)
        ws = wb.create_sheet("Firewall Status", position)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = TAB_COLORS["Firewall Status"]

        sections = cls.parse_firewall_output(firewall_text)
        ws.append(["Profile", "Setting", "Value"])
        if not sections:
            ws.append(["-", "Status",
                       str(firewall_text).strip() if firewall_text else "No data collected."])
            cls.style_data_sheet(ws, "Firewall Status")
            ws.column_dimensions["A"].width = 26
            ws.column_dimensions["C"].width = 46
            return

        for name, rows in sections:
            for offset, (setting, value) in enumerate(rows):
                ws.append([name if offset == 0 else None, setting, value])

        cls.style_data_sheet(ws, "Firewall Status")
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 46

        # Merge the profile name down its block
        row = 2
        for name, rows in sections:
            span = len(rows)
            if span > 1:
                ws.merge_cells(start_row=row, start_column=1, end_row=row + span - 1,
                               end_column=1)
            cls._paint_row(ws, row, 1, 1,
                           font=Font(name="Calibri", size=10, bold=True, color=NAVY),
                           alignment=Alignment(horizontal="left", vertical="center",
                                                wrap_text=True, indent=1))
            for offset in range(span):
                ws.row_dimensions[row + offset].height = 18
            row += span

        ws.freeze_panes = "A2"

    @staticmethod
    def parse_firewall_output(text):
        """
        Turn the netsh dump into [(section_name, [(setting, value), ...]), ...].
        """
        if not text or not isinstance(text, str) or text.strip().upper().startswith("ERROR"):
            return []

        sections = []
        current = None
        for raw_line in text.splitlines():
            line = raw_line.rstrip()
            if not line.strip() or set(line.strip()) == {"-"}:
                continue

            stripped = line.strip()
            if stripped.endswith("Settings:") or stripped == "Logging:":
                current = (stripped.rstrip(":").strip(), [])
                sections.append(current)
                continue

            parts = re.split(r"\s{2,}", stripped, maxsplit=1)
            if len(parts) == 2 and current is not None:
                current[1].append((parts[0].strip(), parts[1].strip()))

        return [(name, rows) for name, rows in sections if rows]

    @classmethod
    def _rebuild_employee_sheet(cls, wb, employee_data):
        """Present the auditor details as a readable two-column list."""
        if "Employee Info" not in wb.sheetnames:
            return

        ws = wb["Employee Info"]
        position = wb.sheetnames.index("Employee Info")
        wb.remove(ws)
        ws = wb.create_sheet("Employee Info", position)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = TAB_COLORS["Employee Info"]

        friendly = {
            "first_name": "First Name",
            "last_name": "Last Name",
            "position": "Position",
            "company_id": "Company ID",
            "company_email": "Company Email",
            "date_hired": "Date Hired",
            "Department": "Department",
        }

        ws.append(["Field", "Value"])
        for key, label in friendly.items():
            if key in employee_data:
                ws.append([label, cls._coerce(str(employee_data[key]))])

        cls.style_data_sheet(ws, "Employee Info")
        ws.column_dimensions["A"].width = 20

        for row_index in range(2, ws.max_row + 1):
            label = ws.cell(row=row_index, column=1)
            label.font = Font(name="Calibri", size=10, bold=True, color=NAVY)
            label.fill = PatternFill(fill_type="solid", fgColor=LABEL_BG)
            label.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # ------------------------------------------------------------------
    # Dashboard
    # ------------------------------------------------------------------
    @classmethod
    def build_dashboard(cls, wb, context):
        """Replace the plain metric table with a full cover page."""
        position = wb.sheetnames.index("Dashboard") if "Dashboard" in wb.sheetnames else 0
        if "Dashboard" in wb.sheetnames:
            wb.remove(wb["Dashboard"])
        ws = wb.create_sheet("Dashboard", position)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = NAVY

        for column, width in zip("ABCDEFGH", (22, 16, 20, 18, 19, 19, 19, 19)):
            ws.column_dimensions[column].width = width

        employee = context["employee"]
        risk_level = context["risk_level"]
        auditor = f"{employee.get('first_name', '')} {employee.get('last_name', '')}".strip()

        # --- Title banner ------------------------------------------------
        cls._merge_block(ws, 1, 1, 8, value="SECURITY AUDIT REPORT",
                         fill=PatternFill(fill_type="solid", fgColor=NAVY),
                         font=Font(name="Calibri", size=20, bold=True, color="FFFFFF"),
                         alignment=Alignment(horizontal="left", vertical="center", indent=1))
        ws.row_dimensions[1].height = 38

        subtitle = f"Endpoint assessment of {context['computer'] or 'this host'}"
        generated = f"Generated {datetime.now():%Y-%m-%d %H:%M} by {auditor or 'unknown auditor'}"
        cls._merge_block(ws, 2, 1, 8, value=f"{subtitle}   \u2022   {generated}",
                         fill=PatternFill(fill_type="solid", fgColor=NAVY_SOFT),
                         font=Font(name="Calibri", size=10, color="D8E4F5"),
                         alignment=Alignment(horizontal="left", vertical="center", indent=1))
        ws.row_dimensions[2].height = 20
        ws.row_dimensions[3].height = 8

        # --- Audit details / security posture ----------------------------
        cls._section_header(ws, 4, 1, 4, "AUDIT DETAILS")
        cls._section_header(ws, 4, 5, 8, "SECURITY POSTURE")

        details = [
            ("Auditor", auditor or "-"),
            ("Position", employee.get("position") or "-"),
            ("Department", employee.get("Department") or employee.get("department") or "-"),
            ("Company ID", employee.get("company_id") or "-"),
            ("Company Email", employee.get("company_email") or "-"),
            ("Computer", context["computer"] or "-"),
            ("Date Hired", employee.get("date_hired") or "-"),
        ]
        for offset, (label, value) in enumerate(details):
            row = 5 + offset
            cls._field(ws, row, label, str(value))

        defender = context["defender"] or {}
        counts = context["record_counts"]
        posture = [
            ("Overall Risk Level", risk_level, cls._risk_style(risk_level)),
            ("Windows Firewall", cls._firewall_summary(wb), None),
            ("Real-time Protection", defender.get("Real-time Protection", "-"), None),
            ("Suspicious Processes", str(context["suspicious_processes"]),
             GREEN if not context["suspicious_processes"] else AMBER),
            ("Local Accounts", str(counts.get("Local Users", 0)), None),
            ("Installed Applications", str(counts.get("Installed Apps List", 0)), None),
            ("USB Devices", str(counts.get("USB History", 0)), None),
        ]
        for offset, (label, value, style) in enumerate(posture):
            row = 5 + offset
            fill, font_color = style or NEUTRAL
            cls._merge_block(
                ws, row, 5, 6, value=label,
                fill=PatternFill(fill_type="solid", fgColor=LABEL_BG),
                font=Font(name="Calibri", size=10, bold=True, color=NAVY),
                alignment=Alignment(horizontal="left", vertical="center", indent=1),
                border=cls._border())
            cls._merge_block(
                ws, row, 7, 8, value=value,
                fill=PatternFill(fill_type="solid", fgColor=fill),
                font=Font(name="Calibri", size=11, bold=True, color=font_color),
                alignment=Alignment(horizontal="left", vertical="center", indent=1),
                border=cls._border())
            ws.row_dimensions[row].height = 19
        ws.row_dimensions[12].height = 8

        # --- KPI cards ---------------------------------------------------
        cls._section_header(ws, 13, 1, 8, "AUDIT SCOPE")
        kpis = [
            ("Installed Apps", counts.get("Installed Apps List", 0)),
            ("Running Processes", counts.get("Running Processes", 0)),
            ("Network Connections", counts.get("Network Connections", 0)),
            ("Scheduled Tasks", counts.get("Scheduled Tasks", 0)),
        ]
        for index, (label, value) in enumerate(kpis):
            first = 1 + index * 2
            cls._merge_block(
                ws, 14, first, first + 1, value=label.upper(),
                fill=PatternFill(fill_type="solid", fgColor=LABEL_BG),
                font=Font(name="Calibri", size=9, bold=True, color=MUTED),
                alignment=Alignment(horizontal="center", vertical="center"),
                border=cls._border())
            cls._merge_block(
                ws, 15, first, first + 1, value=str(value),
                font=Font(name="Calibri", size=22, bold=True, color=NAVY),
                alignment=Alignment(horizontal="center", vertical="center"),
                border=cls._border())
        ws.row_dimensions[14].height = 18
        ws.row_dimensions[15].height = 34
        ws.row_dimensions[16].height = 8

        # --- Collector summary -------------------------------------------
        cls._section_header(ws, 17, 1, 8, "COLLECTOR RESULTS")
        header_font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        header_fill = PatternFill(fill_type="solid", fgColor=LABEL_BG)
        cls._merge_block(ws, 18, 1, 2, value="Collector", fill=header_fill, font=header_font,
                         alignment=Alignment(horizontal="left", vertical="center", indent=1),
                         border=cls._border())
        cls._merge_block(ws, 18, 3, 4, value="Worksheet", fill=header_fill, font=header_font,
                         alignment=Alignment(horizontal="left", vertical="center", indent=1),
                         border=cls._border())
        cls._merge_block(ws, 18, 5, 8, value="Records", fill=header_fill, font=header_font,
                         alignment=Alignment(horizontal="left", vertical="center", indent=1),
                         border=cls._border())
        ws.row_dimensions[18].height = 18

        row = 19
        for index, (label, sheet_name) in enumerate(COLLECTOR_SHEETS):
            count = counts.get(sheet_name, 0)
            fill = PatternFill(fill_type="solid", fgColor=STRIPE) if index % 2 else None
            cls._merge_block(ws, row, 1, 2, value=label, fill=fill,
                             font=Font(name="Calibri", size=10, color=TEXT),
                             alignment=Alignment(horizontal="left", vertical="center", indent=1),
                             border=cls._border())
            cls._merge_block(ws, row, 3, 4, value=sheet_name, fill=fill,
                             font=Font(name="Calibri", size=10, color=MUTED),
                             alignment=Alignment(horizontal="left", vertical="center", indent=1),
                             border=cls._border())
            cls._merge_block(ws, row, 5, 8,
                             value=f"{count:,}" if count else "no data",
                             fill=fill, font=Font(name="Calibri", size=10, color=TEXT),
                             alignment=Alignment(horizontal="left", vertical="center", indent=1),
                             border=cls._border())
            row += 1

        row += 1
        cls._top_processes_block(ws, row, context["top_processes"])
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4)

    @classmethod
    def _top_processes_block(cls, ws, row, top_processes):
        """Top processes table plus a matching bar chart."""
        cls._section_header(ws, row, 1, 8, "TOP 10 RUNNING PROCESSES")

        header_font = Font(name="Calibri", size=10, bold=True, color=NAVY)
        header_fill = PatternFill(fill_type="solid", fgColor=LABEL_BG)
        cls._merge_block(ws, row + 1, 1, 2, value="Process", fill=header_fill,
                         font=header_font,
                         alignment=Alignment(horizontal="left", vertical="center", indent=1),
                         border=cls._border())
        cls._merge_block(ws, row + 1, 3, 4, value="Instances", fill=header_fill,
                         font=header_font,
                         alignment=Alignment(horizontal="left", vertical="center", indent=1),
                         border=cls._border())

        data_start = row + 2
        # Written smallest-first so the busiest process sits at the top of the chart
        entries = sorted(list(top_processes)[:10], key=lambda item: item[1])
        for offset, (name, count) in enumerate(entries):
            fill = PatternFill(fill_type="solid", fgColor=STRIPE) if offset % 2 else None
            cls._merge_block(ws, data_start + offset, 1, 2, value=str(name), fill=fill,
                             font=Font(name="Calibri", size=10, color=TEXT),
                             alignment=Alignment(horizontal="left", vertical="center", indent=1),
                             border=cls._border())
            cls._merge_block(ws, data_start + offset, 3, 4, value=int(count), fill=fill,
                             font=Font(name="Calibri", size=10, color=TEXT),
                             alignment=Alignment(horizontal="left", vertical="center", indent=1),
                             border=cls._border())
            ws.row_dimensions[data_start + offset].height = 17

        if not entries:
            ws.cell(row=data_start, column=1).value = "No process data collected."
            return

        chart = BarChart()
        chart.type = "bar"
        chart.title = "Process instances"
        chart.height = 8.5
        chart.width = 13.5
        chart.legend = None
        chart.gapWidth = 45
        chart.x_axis.delete = False
        chart.y_axis.delete = False

        data = Reference(ws, min_col=3, min_row=data_start - 1,
                         max_row=data_start + len(entries) - 1)
        categories = Reference(ws, min_col=1, min_row=data_start,
                               max_row=data_start + len(entries) - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = True

        chart.series[0].graphicalProperties = GraphicalProperties(solidFill=ACCENT)

        ws.add_chart(chart, f"E{row + 1}")

    # ------------------------------------------------------------------
    # Highlights
    # ------------------------------------------------------------------
    @classmethod
    def apply_highlights(cls, wb):
        """Colour-code findings so they can be scanned quickly."""
        cls._highlight_powershell(wb)
        cls._highlight_processes(wb)
        cls._highlight_users(wb)
        cls._highlight_network(wb)
        cls._highlight_firewall(wb)
        cls._highlight_defender(wb)

    @classmethod
    def _highlight_processes(cls, wb):
        ws = wb["Running Processes"] if "Running Processes" in wb.sheetnames else None
        if ws is None:
            return
        column = cls._column_index(ws, "Name")
        if not column:
            return

        for row in range(2, ws.max_row + 1):
            raw = str(ws.cell(row=row, column=column).value or "")
            # Match on the executable name itself so "hkcmd.exe" is not read as "cmd.exe"
            name = raw.strip().lower().strip('"')
            if name.endswith(".exe"):
                name = name[:-4]
            if name in SUSPICIOUS_PROCESSES:
                cls._fill_row(ws, row, PatternFill(fill_type="solid", fgColor=RED[0]))
                cell = ws.cell(row=row, column=column)
                cell.font = Font(name="Calibri", size=10, bold=True, color=RED[1])

    @classmethod
    def _highlight_firewall(cls, wb):
        """Colour-code the parsed firewall settings (runs after the zebra striping)."""
        ws = wb["Firewall Status"] if "Firewall Status" in wb.sheetnames else None
        if ws is None:
            return

        for row in range(2, ws.max_row + 1):
            setting = str(ws.cell(row=row, column=2).value or "")
            value = str(ws.cell(row=row, column=3).value or "")
            fill, font = cls._state_colour(value, setting)
            cell = ws.cell(row=row, column=3)
            cell.fill = PatternFill(fill_type="solid", fgColor=fill)
            cell.font = Font(name="Calibri", size=10, bold=True, color=font)
            cell.alignment = Alignment(horizontal="left", vertical="center",
                                       wrap_text=True, indent=1)

    @classmethod
    def _highlight_powershell(cls, wb):
        ws = wb["Powershell History"] if "Powershell History" in wb.sheetnames else None
        if ws is None:
            return
        column = cls._column_index(ws, "Suspicious Activity")
        if not column:
            return

        for row in range(2, ws.max_row + 1):
            value = str(ws.cell(row=row, column=column).value or "").strip()
            if value and value.lower() not in ("none", "-", "n/a", "nan"):
                cls._fill_row(ws, row, PatternFill(fill_type="solid", fgColor=RED[0]))
                cell = ws.cell(row=row, column=column)
                cell.font = Font(name="Calibri", size=10, bold=True, color=RED[1])
                cell.alignment = Alignment(horizontal="left", vertical="center")

    @classmethod
    def _highlight_users(cls, wb):
        ws = wb["Local Users"] if "Local Users" in wb.sheetnames else None
        if ws is None:
            return
        for row in range(2, ws.max_row + 1):
            for column in range(1, ws.max_column + 1):
                value = str(ws.cell(row=row, column=column).value or "").strip().lower()
                style = None
                if value == "true":
                    style = GREEN
                elif value == "false":
                    style = RED
                if style:
                    cell = ws.cell(row=row, column=column)
                    cell.fill = PatternFill(fill_type="solid", fgColor=style[0])
                    cell.font = Font(name="Calibri", size=10, bold=True, color=style[1])
                    cell.alignment = Alignment(horizontal="center", vertical="center")

    @classmethod
    def _highlight_network(cls, wb):
        ws = wb["Network Connections"] if "Network Connections" in wb.sheetnames else None
        if ws is None:
            return
        column = cls._column_index(ws, "Status")
        if not column:
            return

        palette = {"ESTABLISHED": AMBER, "LISTEN": NEUTRAL, "LISTENING": NEUTRAL,
                   "SYN_SENT": NEUTRAL, "TIME_WAIT": NEUTRAL, "CLOSE_WAIT": NEUTRAL,
                   "ESTABLISHED ": AMBER}
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=column)
            style = palette.get(str(cell.value or "").strip().upper())
            if style:
                cell.fill = PatternFill(fill_type="solid", fgColor=style[0])
                cell.font = Font(name="Calibri", size=10, bold=True, color=style[1])
                cell.alignment = Alignment(horizontal="center", vertical="center")

    @classmethod
    def _highlight_defender(cls, wb):
        ws = wb["Windows Defender"] if "Windows Defender" in wb.sheetnames else None
        if ws is None:
            return
        for row in range(2, ws.max_row + 1):
            for column in range(1, ws.max_column + 1):
                cell = ws.cell(row=row, column=column)
                text = str(cell.value or "")
                if not text:
                    continue
                style = None
                if "enabled" in text.lower() or "\u2705" in text:
                    style = GREEN
                elif "disabled" in text.lower() or "\u274c" in text:
                    style = RED
                if style:
                    cell.fill = PatternFill(fill_type="solid", fgColor=style[0])
                    cell.font = Font(name="Calibri", size=10, bold=True, color=style[1])
                    cell.alignment = Alignment(horizontal="center", vertical="center")

    # ------------------------------------------------------------------
    # Helpers
    # ------------------------------------------------------------------
    @staticmethod
    def _column_index(ws, header_name):
        for cell in ws[1]:
            if cell.value and str(cell.value).strip().lower() == header_name.lower():
                return cell.column
        return None

    @staticmethod
    def _border():
        side = Side(style="thin", color=BORDER)
        return Border(left=side, right=side, top=side, bottom=side)

    @staticmethod
    def _paint_row(ws, row, first_col, last_col, fill=None, font=None,
                   alignment=None, border=None):
        for column in range(first_col, last_col + 1):
            cell = ws.cell(row=row, column=column)
            if fill is not None:
                cell.fill = fill
            if font is not None:
                cell.font = font
            if alignment is not None:
                cell.alignment = alignment
            if border is not None:
                cell.border = border

    @classmethod
    def _merge_block(cls, ws, row, first_col, last_col, value=None, **styles):
        if last_col > first_col:
            ws.merge_cells(start_row=row, start_column=first_col,
                           end_row=row, end_column=last_col)
        cls._paint_row(ws, row, first_col, last_col, **styles)
        if value is not None:
            ws.cell(row=row, column=first_col).value = value

    @classmethod
    def _section_header(cls, ws, row, first_col, last_col, text):
        cls._merge_block(ws, row, first_col, last_col, value=text,
                         fill=PatternFill(fill_type="solid", fgColor=ACCENT),
                         font=Font(name="Calibri", size=10, bold=True, color="FFFFFF"),
                         alignment=Alignment(horizontal="left", vertical="center", indent=1))
        ws.row_dimensions[row].height = 20

    @classmethod
    def _field(cls, ws, row, label, value):
        cls._merge_block(ws, row, 1, 1, value=label,
                         fill=PatternFill(fill_type="solid", fgColor=LABEL_BG),
                         font=Font(name="Calibri", size=10, bold=True, color=NAVY),
                         alignment=Alignment(horizontal="left", vertical="center", indent=1),
                         border=cls._border())
        cls._merge_block(ws, row, 2, 4, value=value,
                         font=Font(name="Calibri", size=10, color=TEXT),
                         alignment=Alignment(horizontal="left", vertical="center", indent=1),
                         border=cls._border())
        ws.row_dimensions[row].height = 19

    @classmethod
    def _fill_row(cls, ws, row, fill):
        for column in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=column)
            cell.fill = fill
            if cell.alignment is None or cell.alignment.horizontal is None:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    @staticmethod
    def _risk_style(level):
        return RISK_STYLES.get(str(level).upper(), NEUTRAL)

    @staticmethod
    def _state_colour(value, setting=""):
        """Colour a firewall setting by how much it matters, not just its value."""
        text = str(value).strip().lower()
        key = str(setting).strip().lower()

        if key == "state":
            return GREEN if text == "on" else RED
        if key == "firewall policy":
            return GREEN if "blockinbound" in text else AMBER
        if key.startswith("log"):
            return AMBER if text in ("disable", "disabled", "off") else NEUTRAL
        if key == "inboundusernotification":
            return GREEN if text == "enable" else AMBER
        if text in ("off", "disabled", "disable"):
            # Disabling a non-logging setting is the secure default, so stay neutral
            return NEUTRAL
        if text in ("on", "enabled", "enable"):
            return GREEN
        if text.startswith("n/a"):
            return NEUTRAL
        return NEUTRAL

    @staticmethod
    def _firewall_summary(wb):
        """One-line summary of the firewall profile states."""
        if "Firewall Status" not in wb.sheetnames:
            return "-"
        ws = wb["Firewall Status"]

        states = []
        current = None
        for row in range(2, ws.max_row + 1):
            profile = ws.cell(row=row, column=1).value
            if profile:
                current = str(profile)
            setting = ws.cell(row=row, column=2).value
            value = ws.cell(row=row, column=3).value
            if str(setting).strip().lower() == "state" and current:
                states.append((current.replace(" Profile Settings", "").strip(),
                               str(value).strip().upper()))

        if not states:
            return "-"
        enabled = [name for name, state in states if state == "ON"]
        if len(enabled) == len(states):
            return "Enabled (all profiles)"
        if enabled:
            return f"Partial ({len(enabled)}/{len(states)} profiles on)"
        return "Disabled (all profiles)"
